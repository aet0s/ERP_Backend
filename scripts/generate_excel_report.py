import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel_report():
    json_path = os.path.join(os.path.dirname(__file__), 'permission_audit_results.json')
    if not os.path.exists(json_path):
        print(f"Error: {json_path} not found")
        return

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    blue_header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    sub_header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fail_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    accent_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    font_white_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_white_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=12, bold=True, color="1E293B")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="334155")
    font_pass = Font(name="Calibri", size=10, bold=True, color="166534")
    font_fail = Font(name="Calibri", size=10, bold=True, color="991B1B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # ─────────────────────────────────────────────────────────────────────────────
    # SHEET 1: Executive Summary
    # ─────────────────────────────────────────────────────────────────────────────
    ws1 = wb.create_sheet(title="Executive Summary")
    ws1.views.sheetView[0].showGridLines = True

    # Title Block
    ws1.merge_cells("A1:G2")
    title_cell = ws1["A1"]
    title_cell.value = "ERP Role & Permission Matrix Audit Report"
    title_cell.font = font_white_title
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    summary = data.get("summary", {})
    total = summary.get("totalTests", 0)
    passed = summary.get("passed", 0)
    failed = summary.get("failed", 0)
    pass_rate = f"{(passed / total * 100):.1f}%" if total > 0 else "0%"

    meta_rows = [
        ("Execution Timestamp", summary.get("timestamp", "-")),
        ("Target Workspace", summary.get("workspace", "-")),
        ("Owner User Account", summary.get("owner", "-")),
        ("Audited Demo Account", summary.get("demoUser", "-")),
        ("Total Automated Tests", total),
        ("Tests Passed", passed),
        ("Tests Failed", failed),
        ("Overall Success Rate", pass_rate)
    ]

    ws1.cell(row=4, column=1, value="Audit Metadata & Key Metrics").font = font_section

    for idx, (label, val) in enumerate(meta_rows, start=5):
        c1 = ws1.cell(row=idx, column=1, value=label)
        c1.font = font_bold
        c1.fill = sub_header_fill
        c1.border = thin_border

        c2 = ws1.cell(row=idx, column=2, value=str(val))
        c2.font = font_regular
        c2.border = thin_border
        if label == "Tests Passed":
            c2.font = font_pass
            c2.fill = pass_fill
        elif label == "Tests Failed":
            c2.font = font_fail if failed > 0 else font_pass
            c2.fill = fail_fill if failed > 0 else pass_fill
        elif label == "Overall Success Rate":
            c2.font = font_pass
            c2.fill = pass_fill

    # Architectural Fixes Applied
    ws1.cell(row=15, column=1, value="Core Architectural Fixes & Verifications Applied").font = font_section

    fixes = [
        ("1. Same-Person Role Updates", "Fixed: PUT /api/users/:id/role now cleanly handles repeated role changes without unique constraint or master sync collision."),
        ("2. Real-Time Permission Sync", "Fixed: Implemented active session revalidation (5s heartbeat + BroadcastChannel + route-level sync) so member UI reflects owner changes instantly."),
        ("3. Backend API Route Guards", "Fixed: Attached requirePermission(module, 'view') to previously unguarded endpoints (/sales, /sales/:id, /procurements/:id, etc.)."),
        ("4. Page-Level RBAC Protection", "Fixed: Added dynamic usePermissions hook and access restrictions across all 16 page components including AiAnalytics, StockTransfers, and ProductionRuns."),
        ("5. Master DB Synchronization", "Fixed: Synchronized company_users table in Master DB alongside tenant DB so user roles remain consistent across logins."),
        ("6. Administrator Role Coverage", "Fixed: Added 'admin' role to defaultPermissions.js and WORKSPACE_ROLES array ensuring full capability resolution.")
    ]

    headers_fixes = ["Area / Component", "Description of Fix & Resolution"]
    for col_idx, h in enumerate(headers_fixes, start=1):
        cell = ws1.cell(row=17, column=col_idx, value=h)
        cell.font = font_white_header
        cell.fill = blue_header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    for row_idx, (area, desc) in enumerate(fixes, start=18):
        c1 = ws1.cell(row=row_idx, column=1, value=area)
        c1.font = font_bold
        c1.border = thin_border
        c1.fill = accent_fill

        c2 = ws1.cell(row=row_idx, column=2, value=desc)
        c2.font = font_regular
        c2.border = thin_border

    # ─────────────────────────────────────────────────────────────────────────────
    # SHEET 2: Role Assignment Audit
    # ─────────────────────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="Role Assignment Audit")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    t2 = ws2["A1"]
    t2.value = "Role Assignment & Session Verification Audit"
    t2.font = font_white_header
    t2.fill = navy_fill
    t2.alignment = Alignment(horizontal="center", vertical="center")

    role_headers = ["Role Name", "Assignment Status", "Repeated Same-Person Update", "Session Role Match (/auth/me)", "Modules Resolved", "Audit Status", "Detailed Observations"]
    for c_idx, h in enumerate(role_headers, start=1):
        c = ws2.cell(row=3, column=c_idx, value=h)
        c.font = font_white_header
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    for r_idx, r in enumerate(data.get("roleTests", []), start=4):
        vals = [
            r["role"].upper(),
            "YES (200 OK)" if r["assignSuccess"] else "FAILED",
            "YES (Passed)" if r["repeatedChangeSuccess"] else "FAILED",
            "YES (Matched)" if r["authMeRoleMatches"] else "MISMATCH",
            f"{r['permCount']} / 19 Modules",
            r["status"],
            r["details"]
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx == 6:
                cell.font = font_pass if v == "PASS" else font_fail
                cell.fill = pass_fill if v == "PASS" else fail_fill
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 1:
                cell.font = font_bold
                cell.fill = accent_fill

    # ─────────────────────────────────────────────────────────────────────────────
    # SHEET 3: Permission Toggle Tests
    # ─────────────────────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="Permission Toggle Tests")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:I1")
    t3 = ws3["A1"]
    t3.value = "Granular Module Permission Checkbox Toggle Tests (Owner -> Member Reflection)"
    t3.font = font_white_header
    t3.fill = navy_fill
    t3.alignment = Alignment(horizontal="center", vertical="center")

    toggle_headers = ["Role", "Module", "Action Field", "Original", "Toggled", "Owner Save", "Member Reflection (/auth/me)", "API Guard Enforced", "Status", "Test Details"]
    for c_idx, h in enumerate(toggle_headers, start=1):
        c = ws3.cell(row=3, column=c_idx, value=h)
        c.font = font_white_header
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    for r_idx, t in enumerate(data.get("toggleTests", []), start=4):
        vals = [
            t["role"],
            t["module"],
            t["action"],
            t["originalValue"],
            t["toggledValue"],
            "SUCCESS" if t["ownerSaveOk"] else "FAILED",
            "REFLECTED" if t["memberReflectedOk"] else "STALE",
            "ENFORCED (403/200)" if t["apiEnforcedOk"] else "BYPASS DETECTED",
            t["status"],
            t["details"]
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx == 9:
                cell.font = font_pass if v == "PASS" else font_fail
                cell.fill = pass_fill if v == "PASS" else fail_fill
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in [1, 2]:
                cell.font = font_bold
            elif c_idx in [4, 5]:
                cell.alignment = Alignment(horizontal="center")

    # ─────────────────────────────────────────────────────────────────────────────
    # SHEET 4: Page RBAC Guardrails
    # ─────────────────────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="Page RBAC Guardrails")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:G1")
    t4 = ws4["A1"]
    t4.value = "Frontend Page & Route Access Control Guardrails Audit"
    t4.font = font_white_header
    t4.fill = navy_fill
    t4.alignment = Alignment(horizontal="center", vertical="center")

    page_headers = ["Page Component", "Module Key", "URL Route", "usePermissions Hook", "canView Guard Present", "Status", "Security Audit Verification"]
    for c_idx, h in enumerate(page_headers, start=1):
        c = ws4.cell(row=3, column=c_idx, value=h)
        c.font = font_white_header
        c.fill = blue_header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    for r_idx, p in enumerate(data.get("routeAudits", []), start=4):
        vals = [
            p["page"],
            p["module"],
            p["route"],
            "YES (Present)" if p["hasUsePermissions"] else "NO",
            "YES (Enforced)" if p["hasViewGuard"] else "NO",
            p["status"],
            p["details"]
        ]
        for c_idx, v in enumerate(vals, start=1):
            cell = ws4.cell(row=r_idx, column=c_idx, value=v)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx == 6:
                cell.font = font_pass if v == "PASS" else font_fail
                cell.fill = pass_fill if v == "PASS" else fail_fill
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 1:
                cell.font = font_bold

    # Auto-fit column widths across all sheets
    for ws in [ws1, ws2, ws3, ws4]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if cell.row in [1, 2]:
                    continue
                max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output_path = os.path.join(os.path.dirname(__file__), '../../Role_and_Permission_Matrix_Audit_Report.xlsx')
    output_path = os.path.abspath(output_path)
    wb.save(output_path)
    print(f"\nExcel report successfully generated: {output_path}")

if __name__ == '__main__':
    build_excel_report()
